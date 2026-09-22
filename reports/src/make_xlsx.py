"""Build Ket_qua_benchmark_2026-09-16.xlsx from all_results.csv."""

import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

S = sys.argv[1]
df = pd.read_csv(f"{S}/all_results.csv")
NAME = {
    "arima": "ARIMA",
    "xgboost": "XGBoost",
    "random_forest": "Random Forest",
    "gru": "GRU v2",
}
MODELS = ["arima", "xgboost", "random_forest", "gru"]
ORDER = [
    "ACB 1d",
    "FPT 1d",
    "HPG 1d",
    "VCB 1d",
    "VNM 1d",
    "BTCUSDT 1d",
    "ETHUSDT 1d",
    "SOLUSDT 1d",
    "BTCUSDT 1h",
]
df["dataset"] = df["symbol"] + " " + df["timeframe"]
df["order"] = df["dataset"].map({k: i for i, k in enumerate(ORDER)})
df["model_order"] = df["model"].map({m: i for i, m in enumerate(MODELS)})
df = df.sort_values(["order", "model_order"]).reset_index(drop=True)

FONT = "Arial"
base = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
head_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F3864")
good_fill = PatternFill("solid", fgColor="E2F0D9")
bad_fill = PatternFill("solid", fgColor="FBE5D6")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def header(ws, row, labels, widths=None):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def style_body(ws, first_row, last_row, last_col):
    for r in range(first_row, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if cell.font != bold:
                cell.font = base


# ── Sheet 1: raw results, one row per MLflow run ─────────────
ws = wb.active
ws.title = "Ket qua chi tiet"
cols = [
    "Mã",
    "Khung",
    "Mô hình",
    "n test",
    "MAE",
    "RMSE",
    "MAPE (%)",
    "Directional Acc.",
    "Naive MAE",
    "Naive RMSE",
    "Naive MAPE (%)",
    "Δ RMSE vs Naive (%)",
    "RMSE / Naive RMSE",
    "Xếp hạng trong mã",
    "MLflow run_id",
    "Seed",
    "Trạng thái",
    "Test manifest SHA-256",
]
header(
    ws, 1, cols, [10, 7, 14, 7, 11, 11, 10, 11, 11, 11, 11, 13, 12, 11, 34, 6, 11, 66]
)
n = len(df)
for i, r in df.iterrows():
    row = i + 2
    ws.cell(row=row, column=1, value=r.symbol)
    ws.cell(row=row, column=2, value=r.timeframe)
    ws.cell(row=row, column=3, value=NAME[r.model])
    ws.cell(row=row, column=4, value=int(r.n_samples))
    ws.cell(row=row, column=5, value=float(r.mae))
    ws.cell(row=row, column=6, value=float(r.rmse))
    ws.cell(row=row, column=7, value=float(r.mape_pct))
    ws.cell(row=row, column=8, value=float(r.directional_accuracy))
    ws.cell(row=row, column=9, value=float(r.naive_mae))
    ws.cell(row=row, column=10, value=float(r.naive_rmse))
    ws.cell(row=row, column=11, value=float(r.naive_mape_pct))
    # Derived columns are formulas so the sheet stays live if metrics are edited.
    ws.cell(row=row, column=12, value=f"=(J{row}-F{row})/J{row}*100")
    ws.cell(row=row, column=13, value=f"=F{row}/J{row}")
    ws.cell(
        row=row,
        column=14,
        value=f'=COUNTIFS($A$2:$A${n + 1},A{row},$B$2:$B${n + 1},B{row},$F$2:$F${n + 1},"<"&F{row})+1',
    )
    ws.cell(row=row, column=15, value=r.run_id)
    ws.cell(row=row, column=16, value=int(r.seed))
    ws.cell(row=row, column=17, value=r.status)
    ws.cell(row=row, column=18, value=r.test_manifest_sha256)
    for c in (5, 6, 9, 10):
        ws.cell(row=row, column=c).number_format = "#,##0.0000"
    for c in (7, 11, 12):
        ws.cell(row=row, column=c).number_format = "0.00"
    ws.cell(row=row, column=8).number_format = "0.000"
    ws.cell(row=row, column=13).number_format = "0.000"
style_body(ws, 2, n + 1, len(cols))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
RAW_LAST = n + 1

# ── Sheet 2: pivot summary per dataset ───────────────────────
ws2 = wb.create_sheet("Tong hop theo ma")
ws2["A1"] = (
    "RMSE của từng mô hình so với baseline Naive (SUMIFS từ sheet 'Ket qua chi tiet')"
)
ws2["A1"].font = Font(name=FONT, size=11, bold=True)
labels = (
    ["Mã", "Khung", "n test", "Naive RMSE"]
    + [f"RMSE {NAME[m]}" for m in MODELS]
    + [f"{NAME[m]} / Naive" for m in MODELS]
    + [
        "Mô hình tốt nhất",
        "RMSE tốt nhất",
        "Δ tốt nhất vs Naive (%)",
        "Số mô hình vượt Naive",
    ]
)
header(ws2, 3, labels, [10, 7, 7, 11, 12, 12, 12, 12, 11, 11, 11, 11, 15, 11, 13, 12])
R = "'Ket qua chi tiet'"
for i, ds in enumerate(ORDER):
    sym, tf = ds.split()
    row = i + 4
    ws2.cell(row=row, column=1, value=sym)
    ws2.cell(row=row, column=2, value=tf)
    ws2.cell(
        row=row,
        column=3,
        value=f'=SUMIFS({R}!$D$2:$D${RAW_LAST},{R}!$A$2:$A${RAW_LAST},A{row},{R}!$B$2:$B${RAW_LAST},B{row},{R}!$C$2:$C${RAW_LAST},"ARIMA")',
    )
    ws2.cell(
        row=row,
        column=4,
        value=f'=SUMIFS({R}!$J$2:$J${RAW_LAST},{R}!$A$2:$A${RAW_LAST},A{row},{R}!$B$2:$B${RAW_LAST},B{row},{R}!$C$2:$C${RAW_LAST},"ARIMA")',
    )
    for j, m in enumerate(MODELS):
        col = 5 + j
        ws2.cell(
            row=row,
            column=col,
            value=f'=SUMIFS({R}!$F$2:$F${RAW_LAST},{R}!$A$2:$A${RAW_LAST},A{row},{R}!$B$2:$B${RAW_LAST},B{row},{R}!$C$2:$C${RAW_LAST},"{NAME[m]}")',
        )
        ratio_col = 9 + j
        ws2.cell(
            row=row, column=ratio_col, value=f"={get_column_letter(col)}{row}/$D{row}"
        )
    ws2.cell(
        row=row,
        column=13,
        value=f'=SUBSTITUTE(INDEX($E$3:$H$3,MATCH(MIN(E{row}:H{row}),E{row}:H{row},0)),"RMSE ","")',
    )
    ws2.cell(row=row, column=14, value=f"=MIN(E{row}:H{row})")
    ws2.cell(row=row, column=15, value=f"=(D{row}-N{row})/D{row}*100")
    ws2.cell(row=row, column=16, value=f'=COUNTIF(I{row}:L{row},"<1")')
    for c in (4, 5, 6, 7, 8, 14):
        ws2.cell(row=row, column=c).number_format = "#,##0.0000"
    for c in (9, 10, 11, 12):
        ws2.cell(row=row, column=c).number_format = "0.000"
    ws2.cell(row=row, column=15).number_format = "0.00"
last2 = 3 + len(ORDER)
style_body(ws2, 4, last2, len(labels))
# Conditional fill on the ratio block: green < 1 (beats Naive), red >= 2.
ws2.conditional_formatting.add(
    f"I4:L{last2}", CellIsRule(operator="lessThan", formula=["1"], fill=good_fill)
)
ws2.conditional_formatting.add(
    f"I4:L{last2}",
    CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=bad_fill),
)
# Totals row
tr = last2 + 2
ws2.cell(row=tr, column=1, value="Số lần đứng đầu").font = bold
ws2.cell(row=tr + 1, column=1, value="Số lần vượt Naive").font = bold
for j, m in enumerate(MODELS):
    col = 5 + j
    ws2.cell(row=tr - 1, column=col, value=NAME[m]).font = bold
    ws2.cell(row=tr, column=col, value=f'=COUNTIF($M$4:$M${last2},"{NAME[m]}")')
    ws2.cell(
        row=tr + 1,
        column=col,
        value=f'=COUNTIF({get_column_letter(9 + j)}4:{get_column_letter(9 + j)}{last2},"<1")',
    )
ws2.freeze_panes = "C4"

# ── Sheet 3: ACB vs July ─────────────────────────────────────
ws3 = wb.create_sheet("ACB 1d vs 07-2026")
ws3["A1"] = (
    "Benchmark chính thức ACB 1d: evaluator 23/07/2026 (docs/evidence/ACB_1d) so với chạy lại 16/09/2026"
)
ws3["A1"].font = Font(name=FONT, size=11, bold=True)
labels = [
    "Mô hình",
    "RMSE 07/2026",
    "RMSE 16/09/2026",
    "Chênh lệch RMSE",
    "MAE 07/2026",
    "MAE 16/09/2026",
    "DA 07/2026",
    "DA 16/09/2026",
    "Hạng 07/2026",
    "Hạng 16/09/2026",
    "Ghi chú",
]
header(ws3, 3, labels, [14, 12, 13, 13, 12, 13, 11, 12, 10, 12, 50])
july = {
    "arima": (0.27234869637307557, 0.3882019113976774, 0.5128205128205128, 1),
    "random_forest": (0.5256917774782239, 0.6368350104288949, 0.38461538461538464, 2),
    "xgboost": (0.48262003678541876, 0.6483500635961559, 0.48717948717948717, 3),
    "gru": (0.5393094254762704, 0.7360494889815756, 0.47435897435897434, 4),
}
notes = {
    "arima": "Tái lập đúng — cùng commit pipeline",
    "random_forest": "Tái lập đúng — chênh lệch chỉ do số thực (< 1e-12)",
    "xgboost": "Tái lập đúng — chênh lệch chỉ do số thực (< 1e-15)",
    "gru": "GRU v2 (PR #43): residual head, 8 feature, seq 7 — kiến trúc khác tháng 7",
}
acb = df[(df.symbol == "ACB") & (df.timeframe == "1d")].set_index("model")
for i, m in enumerate(MODELS):
    row = 4 + i
    j = july[m]
    ws3.cell(row=row, column=1, value=NAME[m])
    ws3.cell(row=row, column=2, value=j[1])
    ws3.cell(
        row=row,
        column=3,
        value=f'=SUMIFS({R}!$F$2:$F${RAW_LAST},{R}!$A$2:$A${RAW_LAST},"ACB",{R}!$B$2:$B${RAW_LAST},"1d",{R}!$C$2:$C${RAW_LAST},"{NAME[m]}")',
    )
    ws3.cell(row=row, column=4, value=f"=C{row}-B{row}")
    ws3.cell(row=row, column=5, value=j[0])
    ws3.cell(
        row=row,
        column=6,
        value=f'=SUMIFS({R}!$E$2:$E${RAW_LAST},{R}!$A$2:$A${RAW_LAST},"ACB",{R}!$B$2:$B${RAW_LAST},"1d",{R}!$C$2:$C${RAW_LAST},"{NAME[m]}")',
    )
    ws3.cell(row=row, column=7, value=j[2])
    ws3.cell(
        row=row,
        column=8,
        value=f'=SUMIFS({R}!$H$2:$H${RAW_LAST},{R}!$A$2:$A${RAW_LAST},"ACB",{R}!$B$2:$B${RAW_LAST},"1d",{R}!$C$2:$C${RAW_LAST},"{NAME[m]}")',
    )
    ws3.cell(row=row, column=9, value=j[3])
    ws3.cell(row=row, column=10, value=f"=RANK(C{row},$C$4:$C$7,1)")
    ws3.cell(row=row, column=11, value=notes[m])
    for c in (2, 3, 5, 6):
        ws3.cell(row=row, column=c).number_format = "0.0000000000"
    ws3.cell(row=row, column=4).number_format = "0.000E+00"
    for c in (7, 8):
        ws3.cell(row=row, column=c).number_format = "0.000"
ws3.cell(row=8, column=1, value="Naive").font = bold
ws3.cell(row=8, column=2, value=0.3902169022085419).number_format = "0.0000000000"
ws3.cell(
    row=8,
    column=3,
    value=f'=SUMIFS({R}!$J$2:$J${RAW_LAST},{R}!$A$2:$A${RAW_LAST},"ACB",{R}!$B$2:$B${RAW_LAST},"1d",{R}!$C$2:$C${RAW_LAST},"ARIMA")',
).number_format = "0.0000000000"
ws3.cell(row=8, column=5, value=0.2669230769230771).number_format = "0.0000000000"
ws3.cell(
    row=8,
    column=11,
    value="Baseline predicted_close = current_close; nguồn tháng 7: docs/evidence/ACB_1d/benchmark_overview.csv",
)
style_body(ws3, 4, 8, len(labels))

# ── Sheet 4: notes / provenance ──────────────────────────────
ws4 = wb.create_sheet("Ghi chu")
lines = [
    (
        "Nguồn dữ liệu",
        "artifacts/metrics/<model>/<SYMBOL>_<tf>_<run_id>.csv do từng entrypoint train_*.py sinh ra ngày 16/09/2026; tổng hợp bằng script trong phiên demo.",
    ),
    (
        "Dataset",
        "group_dataset_v1, snapshot ohlcv_full_current (fingerprint 381cd2ee…d6d2), target next_close, horizon 1, split 70/15/15 theo thời gian, seed 42.",
    ),
    (
        "Mã nguồn",
        "github.com/Ennela/NCKH-CryptographicAnalysis, nhánh develop, commit cc4ab56 tại thời điểm train (PR #44/#45 merge sau đó không đổi pipeline train).",
    ),
    (
        "Δ RMSE vs Naive (%)",
        "Công thức: (RMSE_naive − RMSE_model) / RMSE_naive × 100. Dương = tốt hơn Naive.",
    ),
    (
        "RMSE / Naive RMSE",
        "< 1 = tốt hơn Naive (tô xanh); ≥ 2 = kém hơn Naive ít nhất 2 lần (tô đỏ).",
    ),
    (
        "Directional Accuracy",
        "Tỷ lệ [0,1] dự đoán đúng hướng sign(pred − current) so với sign(actual − current); của Naive luôn là hằng số nên không liệt kê.",
    ),
    (
        "Trạng thái",
        "ACB 1d: 'valid' theo evaluator services.training.benchmark (run 5a75b56742ac4ba7861ce27a9c3431ed). Các mã khác: 'preliminary' vì evaluator hiện khóa cứng ACB 1d; manifest hash của 4 mô hình trên mỗi mã đã được kiểm tra trùng nhau.",
    ),
    (
        "Kết quả tháng 7",
        "Sheet 'ACB 1d vs 07-2026' cột 07/2026 lấy từ docs/evidence/ACB_1d/benchmark_overview.csv (evaluator 23/07/2026).",
    ),
    (
        "Ô công thức",
        "Các cột Δ, tỷ số, xếp hạng và toàn bộ sheet tổng hợp là công thức tham chiếu sheet 'Ket qua chi tiet'; sửa số liệu ở đó thì các sheet khác tự cập nhật.",
    ),
]
ws4.column_dimensions["A"].width = 24
ws4.column_dimensions["B"].width = 120
for i, (k, v) in enumerate(lines, 1):
    ws4.cell(row=i, column=1, value=k).font = bold
    assert not v.startswith("="), v
    c = ws4.cell(row=i, column=2, value=v)
    c.font = base
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.calculation.fullCalcOnLoad = True
out = f"{S}/Ket_qua_benchmark_2026-09-16.xlsx"
wb.save(out)
print("saved", out)
