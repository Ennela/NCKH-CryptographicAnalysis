"""Produce the revised journal manuscript from the authors' DOCX.

Applies the October 2026 revision: rewrites boilerplate and report-style
passages, adds Related work, Contributions, the Diebold-Mariano and
stationarity results (read from reports/validations), a Discussion with threats
to validity, declarations and an extended reference list. Figure and table
numbers are recomputed after the edits.

Usage:
    python reports/src/revise_paper.py INPUT.docx OUTPUT.docx \
        --dm reports/validations/dm_test_2026-10-19.csv \
        --stationarity reports/validations/stationarity_2026-10-19.csv
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

import docx
import pandas as pd
from docx.oxml.ns import qn
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MODEL_LABELS = {
    "arima": "ARIMA",
    "xgboost": "XGBoost",
    "random_forest": "Random Forest",
    "gru": "GRU v2",
}
# Paragraphs written by this script, keyed by id(); holding the element keeps the
# lxml proxy (and so its id) stable.
NEW: dict[int, etree._Element] = {}


# --------------------------------------------------------------------------- #
# Low-level paragraph helpers
# --------------------------------------------------------------------------- #
def _clear_text_runs(p: etree._Element) -> None:
    """Remove every run that does not hold a drawing."""
    for run in p.findall(qn("w:r")):
        if run.find(".//" + qn("w:drawing")) is None:
            p.remove(run)
    for hyperlink in p.findall(qn("w:hyperlink")):
        p.remove(hyperlink)


def _add_runs(p: etree._Element, text: str) -> None:
    """Append runs; **x** is bold, *x* is italic."""
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not piece:
            continue
        run = etree.SubElement(p, qn("w:r"))
        if piece.startswith("**"):
            etree.SubElement(etree.SubElement(run, qn("w:rPr")), qn("w:b"))
            piece = piece[2:-2]
        elif piece.startswith("*"):
            etree.SubElement(etree.SubElement(run, qn("w:rPr")), qn("w:i"))
            piece = piece[1:-1]
        node = etree.SubElement(run, qn("w:t"))
        node.text = piece
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def set_text(p: etree._Element, text: str) -> etree._Element:
    """Replace the visible text of a paragraph, keeping its style."""
    _clear_text_runs(p)
    _add_runs(p, text)
    NEW[id(p)] = p
    return p


def clone_after(
    anchor: etree._Element, template: etree._Element, text: str
) -> etree._Element:
    """Insert a copy of template (style, numbering) with new text after anchor."""
    p = copy.deepcopy(template)
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    _add_runs(p, text)
    anchor.addnext(p)
    NEW[id(p)] = p
    return p


def chain(
    anchor: etree._Element, items: list[tuple[etree._Element, str]]
) -> etree._Element:
    """Insert several (template, text) paragraphs in order; return the last."""
    for template, text in items:
        anchor = clone_after(anchor, template, text)
    return anchor


def remove(*elements: etree._Element) -> None:
    """Detach elements from the document."""
    for element in elements:
        element.getparent().remove(element)


def set_cell(table: docx.table.Table, row: int, col: int, text: str) -> None:
    """Replace the text of one cell, keeping the first run's formatting."""
    paragraph = table.rows[row].cells[col].paragraphs[0]
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


# --------------------------------------------------------------------------- #
# Tables in the manuscript's booktabs style
# --------------------------------------------------------------------------- #
def _cell_xml(text: str, header: bool, align: str) -> str:
    border = (
        '<w:tcBorders><w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tcBorders>'
        if header
        else ""
    )
    bold = "<w:b/><w:bCs/>" if header else ""
    safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return (
        f'<w:tc><w:tcPr><w:tcW w:w="0" w:type="auto"/>{border}</w:tcPr>'
        f'<w:p><w:pPr><w:pStyle w:val="Compact"/>'
        f'<w:spacing w:before="30" w:after="30" w:line="240" w:lineRule="auto"/>'
        f'<w:jc w:val="{align}"/></w:pPr>'
        f'<w:r><w:rPr>{bold}<w:sz w:val="15"/></w:rPr>'
        f'<w:t xml:space="preserve">{safe}</w:t></w:r></w:p></w:tc>'
    )


def make_table(rows: list[list[str]], align: list[str]) -> etree._Element:
    """Build a table: first row is the header, top and bottom rules only."""
    xml = [
        f'<w:tbl xmlns:w="{W_NS}"><w:tblPr><w:tblStyle w:val="Table"/>'
        '<w:tblW w:w="5000" w:type="pct"/><w:tblBorders>'
        '<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        '</w:tblBorders><w:tblLook w:val="0020" w:firstRow="1" w:lastRow="0" '
        'w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/></w:tblPr>',
        "<w:tblGrid>" + '<w:gridCol w:w="1000"/>' * len(align) + "</w:tblGrid>",
    ]
    for index, row in enumerate(rows):
        header = index == 0
        prefix = "<w:tr><w:trPr><w:tblHeader/></w:trPr>" if header else "<w:tr>"
        xml.append(
            prefix
            + "".join(_cell_xml(c, header, a) for c, a in zip(row, align))
            + "</w:tr>"
        )
    xml.append("</w:tbl>")
    return etree.fromstring("".join(xml))


def _fmt_p(p: float) -> str:
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def dm_table(dm: pd.DataFrame) -> etree._Element:
    """DM statistic (Holm-adjusted p) per model and dataset, squared loss."""
    header = ["Dataset", "n"] + [MODEL_LABELS[m] for m in MODEL_LABELS]
    rows = [header]
    for (symbol, timeframe), group in dm.groupby(["symbol", "timeframe"], sort=False):
        cells = [f"{symbol} {timeframe}", str(int(group["n"].iloc[0]))]
        for model in MODEL_LABELS:
            r = group.loc[group["model"] == model].iloc[0]
            mark = ""
            if r["p_holm_squared"] < 0.05:
                mark = " †" if r["dm_squared"] < 0 else " *"
            statistic = f"{r['dm_squared']:+.2f}".replace("-", "−")
            cells.append(f"{statistic} ({_fmt_p(r['p_holm_squared'])}){mark}")
        rows.append(cells)
    return make_table(rows, ["left", "right"] + ["right"] * 4)


def stationarity_table(st: pd.DataFrame) -> etree._Element:
    """ADF and KPSS p-values on price level and log return."""

    def kpss(p: float) -> str:
        return "<0.01" if p <= 0.01 else (">0.10" if p >= 0.10 else f"{p:.3f}")

    rows = [
        [
            "Dataset",
            "n",
            "Price: ADF p",
            "Price: KPSS p",
            "Log return: ADF p",
            "Log return: KPSS p",
        ]
    ]
    for (symbol, timeframe), group in st.groupby(["symbol", "timeframe"], sort=False):
        close = group.loc[group["series"] == "close"].iloc[0]
        ret = group.loc[group["series"] == "log_return"].iloc[0]
        rows.append(
            [
                f"{symbol} {timeframe}",
                str(int(close["n"])),
                _fmt_p(close["adf_p"]),
                kpss(close["kpss_p"]),
                _fmt_p(ret["adf_p"]),
                kpss(ret["kpss_p"]),
            ]
        )
    return make_table(rows, ["left", "right", "right", "right", "right", "right"])


def adr_table() -> etree._Element:
    rows = [
        ["ID", "Decision", "Reason", "Accepted cost"],
        [
            "D1",
            "PyTorch is the only deep-learning framework",
            "One skill set to review; reference code (mostly Keras) cannot be pasted in",
            "Examples from the literature must be re-implemented",
        ],
        [
            "D2",
            "market.ohlcv is a TimescaleDB hypertable (7-day chunks, 1-day continuous aggregate)",
            "Queries are time ranges of one symbol; the team already knows PostgreSQL",
            "Extra extension in the database image; some DDL is restricted",
        ],
        [
            "D3",
            "MLflow stores runs, metrics, artifacts and model versions",
            "Tracking UI, artifact reload and versioning without custom code",
            "External service dependency; the ml.* tables remain unimplemented",
        ],
        [
            "D4",
            "Work is split by model, not by technical layer",
            "Four members work in parallel and own one pipeline each",
            "Feature sets differ, so the benchmark compares whole pipelines",
        ],
        [
            "D5",
            "Research data are read only from a SHA-256-fingerprinted snapshot",
            "Live data change hourly; results must be reproducible",
            "The snapshot must be refreshed deliberately",
        ],
    ]
    return make_table(rows, ["left"] * 4)


# --------------------------------------------------------------------------- #
# Numbering
# --------------------------------------------------------------------------- #
def renumber(document: docx.document.Document) -> None:
    """Recompute section, figure, table and listing numbers in document order."""
    body = document.element.body
    paragraphs = [
        docx.text.paragraph.Paragraph(p, document) for p in body.iter(qn("w:p"))
    ]
    # Sections 2-5 of the original move down by one because Related work is new.
    for p in paragraphs:
        if p.style.name.startswith("Heading") and id(p._p) not in NEW:
            match = re.match(r"^(\d+)(\.\S*)\s", p.text)
            if match and 2 <= int(match.group(1)) <= 5:
                set_text(
                    p._p,
                    f"{int(match.group(1)) + 1}{match.group(2)} {p.text[match.end() :]}",
                )
    table_map: dict[str, int] = {}
    counters = {"Table": 0, "Fig.": 0, "Listing": 0}
    for p in paragraphs:
        if p.style.name not in ("Table Caption", "CaptionFig", "CaptionList"):
            continue
        for label in counters:
            match = re.match(rf"^{re.escape(label)} ?(\{{T:\w+\}}|\d+)\.?\s*", p.text)
            if match:
                counters[label] += 1
                key = match.group(1).strip("{}").removeprefix("T:")
                if label == "Table":
                    table_map[key] = counters[label]
                set_text(
                    p._p, f"**{label} {counters[label]}.** {p.text[match.end() :]}"
                )

    def resolve(match: re.Match[str]) -> str:
        return str(table_map[match.group(1) or match.group(2)])

    pattern = re.compile(r"\{T:(\w+)\}|(?<=Table )(\d+)")
    placeholder = re.compile(r"\{T:(\w+)\}()")
    for p in paragraphs:
        # Caption numbers are final already; only resolve placeholders inside them.
        active = placeholder if p.style.name == "Table Caption" else pattern
        # Substitute inside each text node so bold/italic runs survive.
        for node in p._p.iterfind(f"{qn('w:r')}/{qn('w:t')}"):
            if node.text and active.search(node.text):
                node.text = active.sub(resolve, node.text)


# --------------------------------------------------------------------------- #
# Content
# --------------------------------------------------------------------------- #
def revise(
    document: docx.document.Document, dm: pd.DataFrame, st: pd.DataFrame
) -> None:
    P = [p._p for p in document.paragraphs]
    T = [t._tbl for t in document.tables]
    tables = list(document.tables)
    h1, h2 = P[97], P[98]
    first, body, bullet, numbered = P[99], P[101], P[195], P[14]
    tcap, figcap = P[100], P[57]

    worse = dm[(dm["p_holm_squared"] < 0.05) & (dm["dm_squared"] < 0)]
    worse_by_model = worse.groupby("model").size().to_dict()
    min_p_better = dm["p_better_squared"].min()
    n_tests = len(dm)

    # Abstract and keywords --------------------------------------------------
    set_text(
        P[7],
        (
            "Price-forecasting models are usually evaluated offline, while the engineering needed to keep "
            "them retrainable, versioned and verifiable inside a running service is rarely reported. We "
            "describe the design and evaluation of an information system that collects daily and hourly "
            "OHLCV data for 15 stocks listed on the Ho Chi Minh City Stock Exchange and 10 USDT "
            "cryptocurrency pairs from Binance, stores them in TimescaleDB, and serves forecasts from "
            "ARIMA, XGBoost, Random Forest and GRU models through a FastAPI service backed by an MLflow "
            "model registry. Research data are frozen in a SHA-256-fingerprinted snapshot, and a "
            "benchmark validator re-derives the test manifest, reloads each artifact and recomputes every "
            "metric before a result is accepted. Re-training three of the four models in a rebuilt "
            "environment reproduced the earlier metrics to within 1e-12. On nine datasets the best model "
            "had a lower RMSE than the persistence (Naive) forecast on four, but a Diebold–Mariano test "
            f"with the Harvey–Leybourne–Newbold correction found none of the {n_tests} model–dataset "
            f"comparisons significantly better than persistence after Holm adjustment, while {len(worse)} "
            "were significantly worse, mostly tree ensembles trained on price levels. ADF and KPSS tests "
            "show that the price series are non-stationary and their log returns stationary, which "
            "accounts for the extrapolation failure of the tree models. The serving path met a p95 "
            "latency of 0.68 s for uncached requests, and the ingestion pipeline completed 450 of 450 "
            "scheduled jobs after a connection-sharing defect was fixed."
        ),
    )
    set_text(
        P[8],
        (
            "**Keywords:** stock price forecasting; cryptocurrency; MLOps; reproducibility; "
            "Diebold–Mariano test; software architecture"
        ),
    )

    # 1.1 Background --------------------------------------------------------
    set_text(
        P[11],
        (
            "Short-horizon forecasting of stock and cryptocurrency prices has been studied with "
            "statistical models, tree ensembles and recurrent neural networks (Sezer et al. 2020). Much "
            "less is reported about operating such models inside an information system. Data must be "
            "collected continuously from rate-limited third-party APIs, the training set must stay fixed "
            "while the database keeps growing, and the model version that answers a request must be "
            "traceable to its code and data. Studies of production machine-learning systems identify "
            "these concerns as a major source of technical debt (Sculley et al. 2015; Amershi et al. 2019)."
        ),
    )
    set_text(
        P[12],
        (
            "This paper treats the forecasting model as a component of a software system rather than as "
            "the object of study. A team of six students built a web system for Vietnamese stocks and "
            "cryptocurrencies in two-week Scrum sprints, and we evaluate it with four research questions:"
        ),
    )
    remove(P[13])
    set_text(
        P[14],
        (
            "RQ1. How can a market-data pipeline tolerate network failures and API rate limits while "
            "still producing a consistent training set?"
        ),
    )
    set_text(
        P[15],
        (
            "RQ2. How can a trained model be moved into a serving component so that the served version, "
            "its training data and its metrics remain identifiable and reproducible?"
        ),
    )
    set_text(
        P[16],
        (
            "RQ3. Do the integrated models forecast better than a persistence baseline once statistical "
            "significance is taken into account?"
        ),
    )
    clone_after(
        P[16],
        numbered,
        (
            "RQ4. Which defects passed the Agile and continuous-integration process, and how were they "
            "detected?"
        ),
    )

    # 1.2 Contributions ------------------------------------------------------
    set_text(P[17], "1.2. Contributions")
    set_text(P[18], "The paper makes four contributions:")
    chain(
        P[18],
        [
            (
                bullet,
                "A containerised three-service architecture (ingestion, training, inference) with a "
                "hard boundary between operational data and fingerprinted research data (Sections 3–4).",
            ),
            (
                bullet,
                "A benchmark validator with eight gates that recomputes, rather than trusts, the "
                "metrics reported by each training run (Section 5.3).",
            ),
            (
                bullet,
                "An evaluation of four forecasting models against persistence on nine datasets, with "
                "Diebold–Mariano tests under multiple-comparison control, stationarity tests, model "
                "explanations and a feature ablation (Section 7).",
            ),
            (
                bullet,
                "A record of the integration and operations defects found during development, with "
                "their causes and measured effect of the fixes (Sections 5.4, 6.4 and 6.6).",
            ),
        ],
    )
    set_text(
        P[21],
        (
            "**Scope statement.** The system is a research prototype and is not intended for investment "
            "advice. No claim about trading profitability is made."
        ),
    )

    # 1.4 Method -> 1.4, plus Related work section -----------------------------
    set_text(
        P[23],
        (
            "Development followed Scrum with two-week sprints and a three-tier Git branching model "
            "(main ← develop ← feature/*); every change was merged through a reviewed pull request with "
            "automated tests and linting. Team rules, including rules for the use of AI coding "
            "assistants, are kept in one file at the repository root (AGENTS.md) and enforced by "
            "pre-commit hooks and CI."
        ),
    )
    clone_after(
        P[23],
        body,
        (
            "The rest of the paper is organised as follows. Section 2 reviews related work. Section 3 "
            "describes data collection and preparation, Section 4 the architecture, and Section 5 the "
            "integration of the models. Section 6 reports testing and operational measurements, "
            "Section 7 the forecasting experiments, Section 8 discusses the research questions and "
            "threats to validity, and Section 9 concludes."
        ),
    )
    anchor = clone_after(P[24].getprevious(), h1, "2. Related work")
    anchor = chain(
        anchor,
        [
            (h2, "2.1. Forecasting stock and cryptocurrency prices"),
            (
                first,
                "Classical approaches model the price or its differences with ARIMA (Box and Jenkins "
                "1970). Machine-learning studies apply tree ensembles such as Random Forest (Breiman "
                "2001) and gradient boosting (Chen and Guestrin 2016), or recurrent networks such as "
                "LSTM and GRU (Cho et al. 2014; Fischer and Krauss 2018). A review of deep-learning "
                "methods for financial series reports large accuracy gains in many studies but also "
                "heterogeneous baselines and evaluation protocols (Sezer et al. 2020). For "
                "cryptocurrencies, McNally et al. (2018) and Hòa et al. (2024) compare recurrent and "
                "convolutional models on Bitcoin, and Giang (2024) applies a hybrid LSTM–GRU model to "
                "Vietnamese stocks.",
            ),
            (
                body,
                "A recurring weakness in this literature is the choice of baseline. Under weak-form "
                "market efficiency (Fama 1970) the random walk is a strong one-step benchmark, as "
                "Meese and Rogoff (1983) showed for exchange rates. Accuracy differences should also "
                "be tested rather than read from point estimates (Diebold and Mariano 1995; Harvey et "
                "al. 1997), and scale-free error measures should be preferred when series are compared "
                "(Hyndman and Koehler 2006). We therefore report every model relative to the "
                "persistence forecast and test each difference.",
            ),
            (h2, "2.2. Engineering machine-learning systems"),
            (
                first,
                "Sculley et al. (2015) argue that model code is a small part of a production "
                "machine-learning system and that data dependencies and configuration create hidden "
                "technical debt. In a case study at Microsoft, Amershi et al. (2019) found data "
                "management and model evolution to be the stages engineers struggle with most. MLOps "
                "addresses these problems by automating training, versioning and deployment "
                "(Kreuzberger et al. 2023), and tools such as MLflow provide experiment tracking and a "
                "model registry (Zaharia et al. 2018). Kapoor and Narayanan (2023) show that data "
                "leakage is a widespread cause of irreproducible results in machine-learning-based "
                "science; the data-locking boundary (Section 4.3) and the validation gates "
                "(Section 5.3) of our system are designed against this failure.",
            ),
            (h2, "2.3. Agile development of data-intensive systems"),
            (
                first,
                "Empirical studies of agile development document benefits for small teams and open "
                "problems in scaling (Dybå and Dingsøyr 2008; Uludağ et al. 2022), documentation "
                "(Islam et al. 2023) and technical debt (Behutiye et al. 2024). In Vietnam, Hoà (2006) "
                "studied Scrum and Extreme Programming in practice, and Đức and Quang (2024) examined "
                "rapid delivery of digital financial products. Reports that combine an agile student "
                "team, a machine-learning component and a production-style deployment, with measured "
                "outcomes, remain uncommon; this paper provides one such case.",
            ),
        ],
    )

    # Section 2 (now 3): data -------------------------------------------------
    remove(P[34], P[35])  # Listing 1 duplicates Table 3.
    set_text(
        P[52],
        (
            "The cleaning stage (services/ingestion/app/cleaning.py) reads market.ohlcv_raw, applies the "
            "four steps in Table 7 and writes market.ohlcv. Every step uses only observations at or "
            "before the current timestamp. Outliers are detected with Tukey's fences (Tukey 1977) and are "
            "flagged rather than removed, so that the models always see the series as published by the "
            "source."
        ),
    )
    remove(P[54], P[55])  # Listing 2: plain pandas calls, already described by Table 7.
    set_text(
        P[67],
        (
            "Scaling is part of each model's training pipeline and never overwrites the stored data. Each "
            "series is split chronologically without shuffling into training (first 70%), validation "
            "(next 15%) and test (last 15%) sets. Scalers are fitted on the training rows only and then "
            "applied to the other two sets (Table 9), which removes the most common source of look-ahead "
            "leakage in time-series studies (Kapoor and Narayanan 2023)."
        ),
    )
    set_text(
        P[79],
        (
            "Change is the percentage difference between the last and the first close; volatility is "
            "the standard deviation of one-period returns. Daily volatility of the crypto pairs "
            "(2.4–5.1%) is roughly twice that of the stocks (1.5–2.8%)."
        ),
    )
    clone_after(
        P[82],
        first,
        (
            "The dashboard uses three standard indicators with their conventional parameters: simple "
            "moving averages, the Relative Strength Index (Wilder 1978) and MACD (Appel 2005). Table 12 "
            "lists the settings and the rules behind the automatic interpretation panel. The same "
            "functions produce the training features, so the dashboard and the models cannot drift apart."
        ),
    )
    set_text(
        P[90],
        (
            "As an example, the ACB chart (Fig. 8) shows a strong rise from late May 2026 accompanied by "
            "higher volume. The price is above SMA 20 and SMA 20 above SMA 50, RSI 14 is neutral and MACD "
            "has crossed below its signal line, so the panel reports an uptrend with weakening momentum."
        ),
    )
    remove(*P[91:97])  # 2.3.4 Summary repeats the section.

    # Section 3 (now 4): architecture -----------------------------------------
    set_text(
        P[99],
        (
            "The backend is split into three services instead of one application because the parts "
            "differ in load profile, dependencies, failure impact and rate of change (Table 13)."
        ),
    )
    set_text(
        P[101],
        (
            "As a result, a long training run does not slow down the forecast API, the four members "
            "editing training code do not touch the serving service, and the serving image does not "
            "carry Optuna (Akiba et al. 2019) or SHAP."
        ),
    )
    set_text(
        P[107],
        (
            "The pipeline separates operational data from research data. Step 5 exports the cleaned "
            "data to an immutable snapshot with a SHA-256 fingerprint, and every training entrypoint "
            "reads only that snapshot, never the live database. New collection runs therefore cannot "
            "change a training set after an experiment has started."
        ),
    )
    set_text(
        P[113],
        (
            "**Unimplemented write path.** The ml.* tables are designed but not yet written to; model "
            "versions, parameters and metrics are kept in MLflow instead (decision D3 in Table {T:adr}). "
            "Consequently, the inference service skips writing ml.prediction because no matching "
            "ml.model_version row exists."
        ),
    )
    set_text(
        P[116],
        (
            "The rate limiter fails open: if Redis is unavailable, requests are still served. For a "
            "service that exposes public market data, losing availability was judged more harmful than "
            "briefly losing rate limiting. The choice should be revisited if user data are stored."
        ),
    )
    set_text(
        P[120],
        (
            "Table {T:adr} summarises five decisions with long-term impact, following the "
            "Architecture Decision Record format (context, decision, consequences)."
        ),
    )
    caption = clone_after(
        P[120], tcap, "Table {T:adr}. Key architectural decisions and their trade-offs."
    )
    caption.addnext(adr_table())
    remove(*P[121:126], *T[16:21])

    # Section 4 (now 5): integration ------------------------------------------
    set_text(
        P[128],
        (
            "Training a model that scores well offline proved less demanding than keeping it correct "
            "inside a running system. Table 17 contrasts the two settings as they arose in this project."
        ),
    )
    set_text(
        P[138],
        (
            "The validator never accepts the metrics a run reports about itself; it recomputes them from "
            "the stored forecast vector and the locked data. Any silent change in features, data or "
            "artifacts therefore appears as a failed gate instead of a plausible number."
        ),
    )
    set_text(
        P[140],
        (
            "Two defects found during an end-to-end demonstration illustrate failure modes that do not "
            "appear when a model is run in a notebook."
        ),
    )
    set_cell(
        tables[23],
        4,
        1,
        (
            "Mount the mlflow_data volume in the inference service (and in both training and inference "
            "in the production compose file, where it was also missing)"
        ),
    )
    set_cell(
        tables[24],
        4,
        1,
        (
            "Synchronise the feature builder with the training code, add the residual term to the "
            "fallback model class and add four tests that check both sides of the contract"
        ),
    )

    # Section 5 (now 6): testing ------------------------------------------------
    set_text(
        P[145],
        (
            "The repository contains 277 automated tests (pytest) that run in about 11 s. Statement "
            "coverage is 78% overall: above 90% for the inference service, 70–90% for training and "
            "15–40% for ingestion. The low ingestion figure comes from integration tests that point to "
            "localhost instead of the isolated test database and are therefore skipped."
        ),
    )
    remove(P[146], P[147], P[148])
    set_text(
        P[155],
        (
            "Latency of POST /api/v1/predict (ACB, 1 day, 5 steps) was measured end to end on the client "
            "side against a local Docker deployment, with requests spaced below the rate limit. Three "
            "scenarios were run per model: the first call after start-up (artifact loaded from MLflow), "
            "an uncached model run (cache cleared before each of 10 requests) and a cache hit (30 "
            "requests). All 164 requests returned HTTP 200."
        ),
    )
    set_text(
        P[157],
        (
            "All four models meet the target of p95 ≤ 2 s when the model is executed. The first request "
            "after start-up takes 4.1–4.7 s for ARIMA and GRU because the artifact is fetched from "
            "MLflow; this cost is paid once per process. Requests were issued sequentially from one "
            "client, so behaviour under concurrent load is not covered by this measurement."
        ),
    )
    set_text(
        P[163],
        (
            "After the fix, which gives each worker process its own connection pool, the system ran for "
            "21 h 36 min without a container restart and with jobs in every scheduled slot (22/22); "
            "450 of 450 jobs succeeded."
        ),
    )
    set_text(
        P[164],
        (
            "The run was stopped before 24 hours, so the success-rate target is met but the duration "
            "target is not."
        ),
    )
    set_text(
        P[170],
        (
            "Burst traffic: once 60 requests per minute are exceeded for one key and address, the Redis "
            "counter returns 429 Too Many Requests (Table 21)."
        ),
    )
    set_cell(
        tables[27],
        1,
        3,
        (
            "Count per fixed one-minute window. Before the fix 52/164 requests received 429 although "
            "below the threshold; after the fix 0/164. In a burst of 65 requests, 60 were served and 5 "
            "rejected"
        ),
    )

    # Section 6 (now 7): experiments --------------------------------------------
    set_text(P[175], "7. Experimental results")
    set_text(P[176], "7.1. Scope of the evaluation")
    set_text(
        P[177],
        (
            "The evaluation addresses two questions: whether the system produces reproducible and "
            "verifiable results (RQ2), and whether the integrated models forecast better than "
            "persistence once sampling uncertainty is considered (RQ3). Benchmark runs were produced on "
            "16 September 2026 in an environment rebuilt from scratch; the significance and "
            "stationarity analyses (Sections 7.5 and 7.6) were run on 19 October 2026 over the same "
            "registered runs and the same locked data."
        ),
    )
    set_text(P[178], "7.2. Experimental setup")
    set_text(P[180], "7.3. Results on the reference dataset (ACB, 1 day)")
    set_text(
        P[182],
        (
            "Three of the four models, re-trained in a different environment and Docker image, "
            "reproduced the July 2026 metrics to within 1e-12. This shows that the training pipeline is "
            "deterministic given the locked data and the fixed seed, and that the data contract holds."
        ),
    )
    set_text(P[183], "7.4. Extension to nine datasets")
    set_text(
        P[185],
        (
            "Table {T:24}. Summary by dataset (positive Δ means lower RMSE than Naive; point estimates, "
            "see Table {T:dm} for significance)."
        ),
    )
    set_text(
        P[206],
        (
            "Table {T:26}. Test-set RMSE (mean ± standard deviation over 3 seeds). Source: "
            "reports/validations/ablation_gru_features_2026-10-09.csv."
        ),
    )
    set_text(
        P[186],
        (
            "By point estimate, at least one model has a lower RMSE than Naive on 4 of the 9 datasets: "
            "GRU on 4, ARIMA on 2, Random Forest on 1 and XGBoost on none. The margins are small "
            "(0.7–6.4%), so they are tested for significance in Section 7.5."
        ),
    )
    remove(P[187])

    anchor = chain(
        P[186],
        [
            (h2, "7.5. Statistical significance against the Naive baseline"),
            (
                first,
                "For each model and dataset we compare the forecast errors with those of the "
                "persistence forecast on the identical test rows, read from the prediction artifact "
                "of the registered model version. The Diebold–Mariano statistic (Diebold and "
                "Mariano 1995) is computed on the loss differential d_t = L(e_Naive,t) − L(e_model,t) "
                "with the small-sample correction of Harvey et al. (1997) and Student-t reference "
                "distribution, so a positive value favours the model. Squared and absolute losses are "
                f"both tested, and the {n_tests} p-values per loss are adjusted with the Holm "
                "procedure (Holm 1979). The analysis (services/training/significance.py) retrains "
                "nothing; its output is stored in reports/validations/dm_test_2026-10-19.csv.",
            ),
            (
                tcap,
                "Table {T:dm}. Diebold–Mariano statistic (Holm-adjusted two-sided p-value), squared "
                "loss. † significantly worse than Naive at the 5% level; * significantly better.",
            ),
        ],
    )
    anchor.addnext(dm_table(dm))
    anchor = anchor.getnext()
    anchor = chain(
        anchor,
        [
            (
                bullet,
                f"No comparison is significantly better than Naive. Even without adjustment, the "
                f"smallest one-sided p-value in favour of a model is {min_p_better:.3f} (GRU on VCB).",
            ),
            (
                bullet,
                f"{len(worse)} of {n_tests} comparisons are significantly worse: XGBoost on "
                f"{worse_by_model.get('xgboost', 0)} datasets, Random Forest on "
                f"{worse_by_model.get('random_forest', 0)}, GRU on {worse_by_model.get('gru', 0)} "
                "(SOLUSDT 1d and BTCUSDT 1h) and ARIMA on none. Absolute loss gives the same set of "
                "rejections.",
            ),
            (
                bullet,
                "The four point-estimate wins of Section 7.4 (ACB, HPG, VCB, VNM) are therefore "
                "indistinguishable from persistence. With 78 test days per stock the test has little "
                "power to detect differences of 1–3% in MSE, so this result means no detectable "
                "difference, not proven equivalence.",
            ),
            (h2, "7.6. Stationarity of prices and returns"),
            (
                first,
                "To check the explanation given for the tree models, we apply the augmented "
                "Dickey–Fuller test (Dickey and Fuller 1979; null hypothesis: unit root) and the KPSS "
                "test (Kwiatkowski et al. 1992; null hypothesis: level stationarity) to the closing "
                "price and to its log return over the full locked series of each dataset.",
            ),
            (
                tcap,
                "Table {T:adf}. Stationarity tests on the locked series. KPSS p-values are "
                "interpolated from tabulated values and bounded to [0.01, 0.10].",
            ),
        ],
    )
    anchor.addnext(stationarity_table(st))
    close = st[st["series"] == "close"]
    ret = st[st["series"] == "log_return"]
    unit_root = int((close["adf_p"] >= 0.05).sum())
    clone_after(
        anchor.getnext(),
        body,
        (
            f"The ADF test does not reject a unit root in the price level for {unit_root} of 9 datasets "
            f"(the exception is VCB), and KPSS rejects level stationarity for "
            f"{int((close['kpss_p'] < 0.05).sum())} of 9. For log returns ADF "
            f"rejects the unit root in all 9 datasets (p < 0.001) and KPSS does not reject stationarity "
            f"in any ({int((ret['kpss_p'] >= 0.05).sum())}/9). Price levels are thus non-stationary while "
            "returns are stationary, which supports modelling the change in price rather than the level."
        ),
    )

    set_text(P[188], "7.7. Technical observations")
    set_text(P[189], "Three observations follow from Tables {T:24}–{T:adf}:")
    set_text(
        P[190],
        (
            "At daily resolution GRU v2 and ARIMA stay within ±5% of the Naive RMSE. Both model a change "
            "(a residual or a difference) rather than the price level, so their forecasts start from the "
            "persistence value."
        ),
    )
    set_text(
        P[191],
        (
            "XGBoost and Random Forest are 3 to 15 times worse than Naive on series with a strong trend "
            "in the test period (FPT, SOLUSDT, BTCUSDT, ETHUSDT). They predict the non-stationary price "
            "level directly, and regression trees cannot produce values outside the range of targets "
            "seen in training, so the forecast saturates when test prices leave that range. This is a "
            "flaw in the choice of target, not data leakage or an implementation error."
        ),
    )
    set_text(
        P[192],
        (
            "At hourly resolution (2,627 test samples) only ARIMA stays close to Naive, while GRU v2 is "
            "about 21% worse. Its sequence length (7) and moving-average windows (7 and 14) were chosen "
            "for daily data and do not suit the noise at hourly frequency."
        ),
    )
    set_text(
        P[193],
        (
            "These observations point to two changes left for future work: training the tree models on "
            "returns, which Table {T:adf} shows to be stationary, and tuning the GRU separately for the "
            "hourly timeframe."
        ),
    )

    set_text(P[199], "7.8. Model explanation")
    set_text(
        P[200],
        (
            "Each training entrypoint writes explainability/feature_importance.json into its MLflow run "
            "with a method suited to the model (Table 25): SHAP TreeExplainer (Lundberg and Lee 2017) "
            "for the tree models, permutation importance for GRU and the coefficient table for ARIMA. "
            "Producing these artifacts required re-training 27 runs, whose RMSE matched the earlier runs "
            "(no difference for GRU and ARIMA, at most 3.6e-12 for Random Forest owing to the order of "
            "floating-point summation in parallel training)."
        ),
    )
    set_text(P[203], "7.9. Ablation of the GRU features")

    # Move "Limitations of the results" after the ablation as 7.10.
    set_text(P[194], "7.10. Limitations of the results")
    set_text(
        P[196],
        (
            "The Diebold–Mariano test assumes covariance-stationary loss differentials and has low power "
            "on 78 daily test samples; non-significant differences must not be read as equivalence."
        ),
    )
    set_text(
        P[197],
        (
            "The cross-model validation gate is applied to the ACB 1-day benchmark only. For the other "
            "eight datasets the test-sample hashes of the four models were confirmed identical, but the "
            "remaining gates were not run."
        ),
    )
    tail = P[211]
    for element in (P[194], P[195], P[196], P[197], P[198]):
        tail.addnext(element)
        tail = element

    # Section 7 (now 8): discussion ---------------------------------------------
    set_text(P[212], "8. Discussion")
    remove(P[213], P[214], T[33], P[215], P[216], T[34])
    set_text(P[217], "8.1. Answers to the research questions")
    anchor = chain(
        P[217],
        [
            (
                first,
                "**RQ1 (data pipeline).** Idempotent upserts on the key (symbol, timeframe, "
                "timestamp), adapters that normalise every source to one type, and a job log per run "
                "kept the stored series free of duplicates (Table 8). After the connection-sharing "
                "defect was fixed, 450 of 450 scheduled jobs succeeded. The pipeline did not protect "
                "against the loss of a dependency at the package-index level: the quarantine of "
                "vnstock stopped stock collection, and the resulting gap could only be partly "
                "back-filled.",
            ),
            (
                body,
                "**RQ2 (model integration).** Registry naming, the fingerprinted snapshot and the "
                "validation gates made each served model traceable to its data and code; three models "
                "reproduced earlier metrics to within 1e-12 in a new environment. The two integration "
                "incidents (Section 5.4) show that the remaining risk lies at the boundaries between "
                "training and serving: a missing volume mount and a feature contract that changed on "
                "one side only.",
            ),
            (
                body,
                "**RQ3 (forecast quality).** No model beats persistence with statistical significance "
                "on any of the nine datasets, and the tree ensembles trained on price levels are "
                "significantly worse on six. This agrees with the strength of the random-walk "
                "benchmark at one-step horizons and shows why a baseline and a significance test "
                "should be mandatory parts of the benchmark protocol rather than optional extras.",
            ),
            (
                body,
                "**RQ4 (process).** Review and CI did not catch defects that appear only under real "
                "operating conditions: shared database connections under prefork concurrency, a rate "
                "limiter whose window was refreshed on every request, a time-zone error that depended "
                "on the host clock, and the two integration incidents. Each was found by measuring the "
                "running system (Tables {T:5}, {T:20} and {T:21}), which suggests that operational measurements "
                "belong in the definition of done for this kind of system.",
            ),
            (h2, "8.2. Threats to validity"),
            (
                first,
                "**Internal validity.** Each dataset uses one chronological split. XGBoost ran with a "
                "single Optuna trial, Random Forest with fixed parameters and ARIMA with a fixed order "
                "(1,1,1) and the invertibility constraint disabled, which makes some ARIMA estimates "
                "unstable (BTCUSDT 1d, HPG). The results describe these configurations, not the best "
                "achievable performance of each algorithm.",
            ),
            (
                body,
                "**Construct validity.** RMSE, MAE and MAPE measure statistical accuracy, not economic "
                "value; no trading strategy was evaluated. The four pipelines use different feature "
                "sets (decision D4), so the benchmark compares complete pipelines rather than "
                "algorithms on identical input.",
            ),
            (
                body,
                "**External validity.** The study covers one team, 25 assets and about two years of "
                "data, of which nine datasets were benchmarked. Results may differ for other markets, "
                "horizons or longer histories.",
            ),
            (
                body,
                "**Engineering limitations.** Inference always loads the latest registry version, so "
                "the validator does not yet decide which version is served. The API key is exposed to "
                "the browser through a NEXT_PUBLIC_* variable and there is no user login. The "
                "frontend has no automated tests, forecasts are not yet stored in ml.prediction, and "
                "the reliability run lasted 21 h 36 min instead of 24 h.",
            ),
        ],
    )
    remove(*P[218:229])

    # Conclusion -----------------------------------------------------------------
    set_text(P[229], "9. Conclusion")
    set_text(
        P[230],
        (
            "We built and evaluated an information system that collects Vietnamese stock and "
            "cryptocurrency prices, locks them for research, trains and versions four forecasting models "
            "and serves them through a web API. The system met its latency target and, after one "
            "concurrency fix, completed all scheduled ingestion jobs; its training pipeline reproduced "
            "earlier results to within 1e-12 on a rebuilt environment. Forecast quality was the weak "
            "point: no model outperformed persistence with statistical significance, and the tree "
            "ensembles failed on trending series because they were trained on non-stationary price "
            "levels."
        ),
    )
    set_text(
        P[231],
        (
            "Future work will train the tree models on returns, replace the single split with "
            "rolling-origin evaluation (Bergmeir and Benítez 2012), measure latency under concurrent "
            "load, let the validator gate which registry version is served, and restore stock "
            "collection with a source that does not depend on the quarantined package."
        ),
    )
    anchor = chain(
        P[231],
        [
            (h1, "Declaration of generative AI use"),
            (
                first,
                "AI coding assistants were used during software development under the rules in "
                "AGENTS.md, and generative AI tools were used for language editing of this manuscript. "
                "The authors reviewed and verified all code, results and text and take full "
                "responsibility for the content.",
            ),
            (h1, "Data and code availability"),
            (
                first,
                "The source code, the dataset contract, the snapshot fingerprints and the result "
                "files referenced in this paper (reports/validations/) are available at "
                "[repository URL to be added before submission].",
            ),
        ],
    )

    # References -------------------------------------------------------------------
    ref_template = P[243]
    remove(*P[234:])
    anchor = P[233]
    for reference in REFERENCES:
        anchor = clone_after(anchor, ref_template, reference)

    # Broken pandoc captions inside picture paragraphs -----------------------------
    for p in list(document.element.body.iter(qn("w:p"))):
        text = "".join(
            node.text or "" for node in p.iterfind(f"{qn('w:r')}/{qn('w:t')}")
        )
        if "custom-style" not in text:
            continue
        caption_text = re.sub(r'\{?custom-style="CaptionFig"\}?|:::|::', " ", text)
        caption_text = " ".join(caption_text.split())
        clone_after(p, figcap, f"Fig. 0. {caption_text}")
        if p.find(".//" + qn("w:drawing")) is None:
            remove(p)
        else:
            _clear_text_runs(p)


REFERENCES = [
    "Akiba T, Sano S, Yanase T, Ohta T, Koyama M (2019) Optuna: A next-generation hyperparameter "
    "optimization framework. In: Proc. 25th ACM SIGKDD International Conference on Knowledge "
    "Discovery and Data Mining, pp. 2623–2631. https://doi.org/10.1145/3292500.3330701",
    "Amershi S, Begel A, Bird C, DeLine R, Gall H, Kamar E, Nagappan N, Nushi B, Zimmermann T (2019) "
    "Software engineering for machine learning: A case study. In: Proc. 41st International "
    "Conference on Software Engineering: Software Engineering in Practice (ICSE-SEIP), pp. 291–300. "
    "https://doi.org/10.1109/ICSE-SEIP.2019.00042",
    "Appel G (2005) Technical Analysis: Power Tools for Active Investors. FT Prentice Hall, Upper "
    "Saddle River, NJ.",
    "Behutiye WN, Rodriguez P, Oivo M, Tosun A (2024) Analyzing the concept of technical debt in the "
    "context of agile software development: A systematic literature review. arXiv preprint "
    "arXiv:2401.14882. https://doi.org/10.48550/arXiv.2401.14882",
    "Bergmeir C, Benítez JM (2012) On the use of cross-validation for time series predictor "
    "evaluation. Information Sciences 191:192–213. https://doi.org/10.1016/j.ins.2011.12.028",
    "Box GEP, Jenkins GM (1970) Time Series Analysis: Forecasting and Control. Holden-Day, San "
    "Francisco.",
    "Breiman L (2001) Random forests. Machine Learning 45(1):5–32. "
    "https://doi.org/10.1023/A:1010933404324",
    "Chen T, Guestrin C (2016) XGBoost: A scalable tree boosting system. In: Proc. 22nd ACM SIGKDD "
    "International Conference on Knowledge Discovery and Data Mining, pp. 785–794. "
    "https://doi.org/10.1145/2939672.2939785",
    "Cho K, van Merriënboer B, Gulcehre C, Bahdanau D, Bougares F, Schwenk H, Bengio Y (2014) "
    "Learning phrase representations using RNN encoder–decoder for statistical machine "
    "translation. In: Proc. 2014 Conference on Empirical Methods in Natural Language Processing "
    "(EMNLP), pp. 1724–1734. https://doi.org/10.3115/v1/D14-1179",
    "Dickey DA, Fuller WA (1979) Distribution of the estimators for autoregressive time series with "
    "a unit root. Journal of the American Statistical Association 74(366):427–431. "
    "https://doi.org/10.2307/2286348",
    "Diebold FX, Mariano RS (1995) Comparing predictive accuracy. Journal of Business & Economic "
    "Statistics 13(3):253–263. https://doi.org/10.1080/07350015.1995.10524599",
    "Đức PT, Quang NN (2024) Rapid development and delivery of digital financial products and "
    "services using low-code platforms. Journal of Economics, Law & Banking 26(266):41–50.",
    "Dybå T, Dingsøyr T (2008) Empirical studies of agile software development: A systematic review. "
    "Information and Software Technology 50(9–10):833–859. "
    "https://doi.org/10.1016/j.infsof.2008.01.006",
    "Fama EF (1970) Efficient capital markets: A review of theory and empirical work. The Journal of "
    "Finance 25(2):383–417. https://doi.org/10.2307/2325486",
    "Fischer T, Krauss C (2018) Deep learning with long short-term memory networks for financial "
    "market predictions. European Journal of Operational Research 270(2):654–669. "
    "https://doi.org/10.1016/j.ejor.2017.11.054",
    "Giang TV (2024) Stock price forecasting using a hybrid LSTM-GRU model. Finance Magazine, "
    "October, pp. 59–62.",
    "Harvey D, Leybourne S, Newbold P (1997) Testing the equality of prediction mean squared errors. "
    "International Journal of Forecasting 13(2):281–291. "
    "https://doi.org/10.1016/S0169-2070(96)00719-4",
    "Hoà PQ (2006) Software development applying Scrum and Extreme Programming methods. Master's "
    "thesis.",
    "Hòa TT, Thạnh LM, Cương NĐH (2024) Comparing the performance of 1D-CNN, LSTM, and hybrid deep "
    "learning models in forecasting Bitcoin closing prices. Journal of Science and Technology, "
    "University of Sciences, Hue University 27(1B):121–132.",
    "Holm S (1979) A simple sequentially rejective multiple test procedure. Scandinavian Journal of "
    "Statistics 6(2):65–70.",
    "Hyndman RJ, Koehler AB (2006) Another look at measures of forecast accuracy. International "
    "Journal of Forecasting 22(4):679–688. https://doi.org/10.1016/j.ijforecast.2006.03.001",
    "Islam MA, Hasan R, Eisty NU (2023) Documentation practices in agile software development: A "
    "systematic literature review. In: Proc. 21st IEEE/ACIS International Conference on Software "
    "Engineering, Management and Applications (SERA), Orlando, FL, USA. "
    "https://doi.org/10.48550/arXiv.2304.07482",
    "Kapoor S, Narayanan A (2023) Leakage and the reproducibility crisis in machine-learning-based "
    "science. Patterns 4(9):100804. https://doi.org/10.1016/j.patter.2023.100804",
    "Kreuzberger D, Kühl N, Hirschl S (2023) Machine learning operations (MLOps): Overview, "
    "definition, and architecture. IEEE Access 11:31866–31879. "
    "https://doi.org/10.1109/ACCESS.2023.3262138",
    "Kwiatkowski D, Phillips PCB, Schmidt P, Shin Y (1992) Testing the null hypothesis of "
    "stationarity against the alternative of a unit root. Journal of Econometrics 54(1–3):159–178. "
    "https://doi.org/10.1016/0304-4076(92)90104-Y",
    "Lundberg SM, Lee S-I (2017) A unified approach to interpreting model predictions. In: Advances "
    "in Neural Information Processing Systems 30, pp. 4765–4774.",
    "McNally S, Roche J, Caton S (2018) Predicting the price of Bitcoin using machine learning. In: "
    "Proc. 26th Euromicro International Conference on Parallel, Distributed and Network-based "
    "Processing (PDP), pp. 339–343. https://doi.org/10.1109/PDP2018.2018.00060",
    "Meese RA, Rogoff K (1983) Empirical exchange rate models of the seventies: Do they fit out of "
    "sample? Journal of International Economics 14(1–2):3–24. "
    "https://doi.org/10.1016/0022-1996(83)90017-X",
    "Sculley D, Holt G, Golovin D, Davydov E, Phillips T, Ebner D, Chaudhary V, Young M, Crespo J-F, "
    "Dennison D (2015) Hidden technical debt in machine learning systems. In: Advances in Neural "
    "Information Processing Systems 28, pp. 2503–2511.",
    "Sezer OB, Gudelek MU, Ozbayoglu AM (2020) Financial time series forecasting with deep learning: "
    "A systematic literature review: 2005–2019. Applied Soft Computing 90:106181. "
    "https://doi.org/10.1016/j.asoc.2020.106181",
    "Tukey JW (1977) Exploratory Data Analysis. Addison-Wesley, Reading, MA.",
    "Uludağ Ö, Philipp P, Putta A, Paasivaara M, Lassenius C, Matthes F (2022) Revealing the state of "
    "the art of large-scale agile development research: A systematic mapping study. Journal of "
    "Systems and Software 194:111473. https://doi.org/10.1016/j.jss.2022.111473",
    "Wilder JW (1978) New Concepts in Technical Trading Systems. Trend Research, Greensboro, NC.",
    "Zaharia M, Chen A, Davidson A, Ghodsi A, Hong SA, Konwinski A, Murching S, Nykodym T, Ogilvie P, "
    "Parkhe M, Xie F, Zumar C (2018) Accelerating the machine learning lifecycle with MLflow. IEEE "
    "Data Engineering Bulletin 41(4):39–45.",
]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--dm", type=Path, required=True)
    parser.add_argument("--stationarity", type=Path, required=True)
    args = parser.parse_args()
    document = docx.Document(str(args.input))
    revise(document, pd.read_csv(args.dm), pd.read_csv(args.stationarity))
    renumber(document)
    document.save(str(args.output))


if __name__ == "__main__":
    main()
